"""The corrected sheet: everything the brain learned, written back as plain cells a person can review.

* rows in the order they actually ran
* gaps replaced by explicit references ({{Login.access_token}}, {{$email}}) so the next run is repeatable
* fixes from healing (added fields and headers, corrected methods and values) applied
* blank ExpectedResponse cells filled with the real response (values that change on every call become "*")
* a RestoraNotes column saying what changed and why
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from . import jsonpath
from .heal import set_query
from .names import norm
from .volatile import looks_volatile

COLUMNS = ["Name", "Method", "URL", "Headers", "Body", "ExpectedStatus", "ExpectedResponse", "RestoraNotes"]


def _wildcard(value, key="", per_run=frozenset()):
    """Replace values that change on every call with "*" (for a captured sample). per_run: values Restora
    generated or copied in this run, which will be different next time."""
    if isinstance(value, dict):
        return {k: _wildcard(v, k, per_run) for k, v in value.items()}
    if isinstance(value, list):
        return [_wildcard(v, key, per_run) for v in value]
    from .volatile import why_dynamic
    if str(value) in per_run or looks_volatile(key, value) or why_dynamic(key, value, _bump(value)) is not None:
        return "*"
    return value


def _bump(v):
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, (int, float)):
        return v + 1
    return str(v) + "0"


def build_row(record, capture: bool = True) -> dict:
    fix = record.fix or {}
    row = fix.get("row")
    if row is None:
        return {}
    notes = list(record.planned)
    method = (record.method or row.method)
    url = row.url
    headers = dict(row.headers)
    body_text = row.body
    data = None
    if body_text and body_text.lstrip()[:1] in "{[":
        try:
            data = json.loads(body_text)
        except ValueError:
            data = None

    # 1. gaps -> explicit references (or the literal value when nothing better exists)
    for f in fix.get("fills", []):
        ph = f.placeholder
        if ph.explicit and f.ref == ph.raw:
            continue                                       # already explicit
        replacement = f.ref or (None if f.source == "memory" else f.value)
        if replacement is None:
            continue
        if ph.where in ("url", "query"):
            url = url.replace(ph.raw, str(replacement), 1)
        elif ph.where == "header":
            k = ph.key
            current = headers.get(k, "")
            headers[k] = current.replace(ph.raw, str(replacement), 1) if not ph.whole else (
                str(replacement) if not (k.lower() == "authorization" and " " not in str(replacement)) else "Bearer " + str(replacement))
        elif ph.where == "body":
            if data is not None and ph.path:
                if ph.whole:
                    jsonpath.put(data, ph.path, replacement)
                else:
                    cur = jsonpath.get(data, ph.path)
                    jsonpath.put(data, ph.path, str(cur).replace(ph.raw, str(replacement), 1))
            elif body_text:
                body_text = body_text.replace(ph.raw, str(replacement), 1)
        notes.append(f"filled {ph.label.split(' = ')[0]} ({f.reason})")

    # 2. fixes that healing proved to work
    for patch in fix.get("patches", []):
        for ch in patch.changes:
            value = ch.ref or ch.value
            if ch.where == "method":
                method = str(ch.value)
            elif ch.where == "header":
                key = next((k for k in headers if k.lower() == ch.name.lower()), ch.name)
                if ch.action == "remove":
                    headers.pop(key, None)
                else:
                    headers[key] = str(value)
            elif ch.where == "query":
                if ch.ref:
                    sep = "&" if "?" in url else "?"
                    url = f"{url}{sep}{ch.name}={ch.ref}"
                else:
                    url = set_query(url, ch.name, None if ch.action == "remove" else ch.value)
            elif ch.where == "body":
                if ch.name == "$":
                    data = ch.value
                    continue
                if data is None:
                    data = {}
                parts = jsonpath.parse(ch.name)
                if ch.action == "remove":
                    jsonpath.delete(data, parts)
                else:
                    jsonpath.put(data, parts, value)
            if ch.note and ch.where != "wait":
                notes.append(ch.note)

    if data is not None:
        body_text = json.dumps(data, ensure_ascii=False)

    # 3. capture the real response when the sheet had no sample
    expected_status = row.original.get("expected_status", "") if row.original else ""
    expected_response = row.original.get("expected_response", "") if row.original else ""
    if capture and not expected_response and record.ok and fix.get("outcome") is not None:
        out = fix["outcome"]
        per_run = {str(f.value) for f in fix.get("fills", []) if f.source in ("generated", "memory", "ref")}
        per_run |= {str(c.value) for p in fix.get("patches", []) for c in p.changes if c.ref}
        per_run |= set(fix.get("generated") or ())
        if out.is_json:
            expected_response = json.dumps(_wildcard(out.json_value, "", frozenset(per_run)), ensure_ascii=False)
        elif out.text and len(out.text) <= 2000:
            expected_response = out.text.strip()
        if expected_response:
            expected_status = expected_status or str(out.status)
            notes.append("ExpectedResponse captured from the real response (values that change each call set to *) - please review")
    elif row.expected_response and row.original and row.expected_response != row.original.get("expected_response"):
        expected_response = row.expected_response
    dynamic = fix.get("dynamic") or []
    if dynamic and expected_response and expected_response.lstrip()[:1] in "{[":
        try:
            sample = json.loads(expected_response)
            for path in dynamic:
                parts = jsonpath.parse(path)
                if parts and jsonpath.exists(sample, parts):
                    jsonpath.put(sample, parts, "*")
            expected_response = json.dumps(sample, ensure_ascii=False)
            names = ", ".join(dict.fromkeys(p.split(".")[-1] for p in dynamic))
            notes.append(f"marked values that change on every call with * in ExpectedResponse ({names})")
        except ValueError:
            pass
    if record.result == "SKIPPED" and row.original:
        o = row.original
        return {"Name": o.get("name", ""), "Method": o.get("method", ""), "URL": o.get("url", ""),
                "Headers": o.get("headers", ""), "Body": o.get("body", ""), "ExpectedStatus": o.get("expected_status", ""),
                "ExpectedResponse": o.get("expected_response", ""), "RestoraNotes": "switched off (#) - left as it was"}
    return {
        "Name": (row.original.get("name") or row.name) if row.original else row.name,
        "Method": method,
        "URL": url,
        "Headers": "\n".join(f"{k}: {v}" for k, v in headers.items()),
        "Body": body_text or "",
        "ExpectedStatus": expected_status,
        "ExpectedResponse": expected_response,
        "RestoraNotes": "\n".join(dict.fromkeys(n for n in notes if n)),
    }


def build(records: list, capture: bool = True) -> list:
    rows = []
    for r in sorted(records, key=lambda x: x.position or x.number):
        built = build_row(r, capture)
        if built:
            rows.append(built)
    return rows


def write_csv(rows: list, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow(r)
    return path


def write_xlsx(rows: list, path: Path) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Tests (corrected)"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B4C7E")
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])
    widths = {"A": 34, "B": 9, "C": 52, "D": 40, "E": 56, "F": 14, "G": 56, "H": 70}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def same_name(a: str, b: str) -> bool:
    return norm(a) == norm(b)
