"""Read the input sheet (CSV or Excel) into Row objects.

The sheet only needs two things per row: a URL and (optionally) a Method.
Everything else is optional and has a sensible default - see README.md.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

from .status import parse_status
from .text import loads_lenient, looks_like_json


class InputError(Exception):
    """The input file itself is unusable (missing, no URL column, ...). Stops the run."""


METHODS = {"GET", "POST", "PUT", "PATCH", "DELETE", "HEAD", "OPTIONS"}

# Column names are matched ignoring case, spaces and punctuation, so "Request Body" == "request_body".
ALIASES = {
    "name": ["name", "testname", "testcase", "title", "description", "apiname", "scenario"],
    "method": ["method", "type", "requesttype", "requestmethod", "httpmethod", "verb"],
    "url": ["url", "apiurl", "endpoint", "requesturl", "link", "uri", "api"],
    "headers": ["headers", "header", "requestheaders"],
    "body": ["body", "requestbody", "payload", "request", "data", "input"],
    "expected_status": ["expectedstatus", "status", "statuscode", "expectedstatuscode", "httpstatus", "expectedcode"],
    "expected_response": ["expectedresponse", "sampleresponse", "response", "expectedbody",
                          "expected", "expectedoutput", "output"],
}
CANONICAL = {  # headings used when we write a CSV (Postman import)
    "name": "Name", "method": "Method", "url": "URL", "headers": "Headers", "body": "Body",
    "expected_status": "ExpectedStatus", "expected_response": "ExpectedResponse",
}


@dataclass
class Row:
    number: int                     # 1, 2, 3 ... in sheet order (used in test ids and reports)
    sheet_row: int                  # the row number you see in Excel
    name: str
    method: str
    url: str
    headers: dict = field(default_factory=dict)
    body: str | None = None
    expected_status: int | None = None
    expected_response: str | None = None
    skip: str | None = None         # reason, when the row is commented out
    error: str | None = None        # reason, when the row is unusable (the row fails; the run continues)
    notes: list = field(default_factory=list)   # things we auto-fixed / guessed, shown in the report
    problems: dict = field(default_factory=dict)  # code -> message; the brain may fix some (url, body)
    method_inferred: bool = False   # Method was blank in the sheet
    original: dict = field(default_factory=dict)  # the row exactly as typed (for the corrected sheet)
    brain: dict = field(default_factory=dict)     # what the brain planned for this row (see restora.brain)

    def refresh_error(self) -> None:
        self.error = "; ".join(self.problems.values()) if self.problems and not self.skip else None


# --------------------------------------------------------------------------- reading files

def _norm(heading: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(heading).lower())


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_csv(path: Path) -> list[list[str]]:
    text = _decode(path.read_bytes())
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    delimiter = max([",", ";", "\t"], key=first.count) if first else ","
    return [list(r) for r in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)]


def _read_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise InputError("Reading .xlsx needs the 'openpyxl' package (pip install openpyxl), "
                         "or save the sheet as CSV.") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out = []
    for cells in ws.iter_rows(values_only=True):
        row = []
        for v in cells:
            if v is None:
                row.append("")
            elif isinstance(v, float) and v.is_integer():
                row.append(str(int(v)))
            else:
                row.append(str(v))
        out.append(row)
    wb.close()
    return out


def read_table(path: Path) -> list[list[str]]:
    if not path.exists():
        raise InputError(f"Input file not found: {path}")
    ext = path.suffix.lower()
    if ext == ".csv" or ext == ".txt":
        return _read_csv(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise InputError(f"Unsupported input type '{ext}'. Use .csv, .xlsx, or a Postman collection (.json).")


# --------------------------------------------------------------------------- cell parsers

def parse_headers(text: str) -> dict:
    """'Name: value' lines (one per line), or a JSON object."""
    text = (text or "").strip()
    if not text:
        return {}
    if text.startswith("{"):
        data, _, _ = loads_lenient(text)
        if not isinstance(data, dict):
            raise ValueError("Headers JSON must be an object like {\"Name\": \"value\"}")
        return {str(k): str(v) for k, v in data.items()}
    headers = {}
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line:
            continue
        if ":" not in line:
            raise ValueError(f"Header line '{line}' must look like 'Name: value' (one header per line)")
        k, v = line.split(":", 1)
        headers[k.strip()] = v.strip()
    return headers


def _parse_status(text: str) -> int | None:
    return parse_status(text)                      # 201, or a class such as 4xx (see status.py)


def _auto_name(method: str, url: str) -> str:
    p = urlparse(url)
    where = (p.path or "/") + (f"?{p.query}" if p.query else "")
    return f"{method} {where}"[:70]


# --------------------------------------------------------------------------- building rows

def _build_row(number: int, sheet_row: int, cells: dict) -> Row:
    name = cells.get("name", "").strip()
    method_raw = cells.get("method", "").strip()
    url = cells.get("url", "").strip()
    body = cells.get("body", "")
    row = Row(number=number, sheet_row=sheet_row, name=name, method=method_raw.upper(), url=url)

    # commented-out rows: a '#' at the start of Name, Method or URL
    for label, val in (("Name", name), ("Method", method_raw), ("URL", url)):
        if val.startswith("#"):
            row.skip = f"commented out with # in {label}"
            row.method = method_raw.lstrip("#").strip().upper()
            row.name = name.lstrip("#").strip()
            row.url = url.lstrip("#").strip()
            break

    problems: dict = {}
    body_text = body if body.strip() else None
    row.original = {"name": name, "method": method_raw, "url": url, "headers": cells.get("headers", ""), "body": body,
                    "expected_status": cells.get("expected_status", ""), "expected_response": cells.get("expected_response", "")}
    if not row.method:
        row.method = "POST" if body_text else "GET"
        row.method_inferred = True
        row.notes.append(f"Method was blank, so {row.method} was used ({'there is a body' if body_text else 'no body'}).")
    if row.method not in METHODS:
        problems["method"] = f"Method '{row.method}' is not valid (use GET, POST, PUT, PATCH or DELETE)"
    if not row.url:
        problems["url"] = "URL is empty"
    elif not re.match(r"^https?://", row.url, re.I):
        problems["url"] = f"URL '{row.url}' must start with http:// or https://"
    if not row.name:
        row.name = _auto_name(row.method, row.url)

    try:
        row.headers = parse_headers(cells.get("headers", ""))
    except ValueError as e:
        problems["headers"] = f"Headers: {e}"

    if body_text is not None:
        if looks_like_json(body_text):
            try:
                _, fixed, was_fixed = loads_lenient(body_text)
                if was_fixed:
                    body_text = fixed
                    row.notes.append("Body contained curly 'smart quotes'; they were replaced with normal quotes.")
            except ValueError as e:
                row.notes.append(f"Body looks like JSON but is not valid JSON ({e}). It was sent exactly as typed.")
                problems["body_json"] = ""          # not an error by itself; the brain can repair it
        row.body = body_text

    try:
        row.expected_status = _parse_status(cells.get("expected_status", ""))
    except ValueError as e:
        problems["status"] = str(e)

    exp = cells.get("expected_response", "")
    if exp.strip():
        if looks_like_json(exp):
            try:
                _, fixed, was_fixed = loads_lenient(exp)
                exp = fixed
                if was_fixed:
                    row.notes.append("ExpectedResponse contained smart quotes; they were replaced.")
            except ValueError as e:
                problems["expected_json"] = f"ExpectedResponse is not valid JSON ({e})"
        row.expected_response = exp.strip()

    row.problems = {k: v for k, v in problems.items() if v}
    if "body_json" in problems:
        row.brain["body_invalid"] = True
    row.refresh_error()
    return row


@dataclass
class Sheet:
    rows: list
    ignored_columns: list


def load_rows(path: str | Path) -> list[Row]:
    return load_sheet(path).rows


def load_sheet(path: str | Path) -> Sheet:
    path = Path(path)
    table = read_table(path)
    header_idx = next((i for i, r in enumerate(table) if any(c.strip() for c in r)), None)
    if header_idx is None:
        raise InputError(f"{path.name} is empty.")

    column_of: dict[str, int] = {}
    unknown = []
    for i, heading in enumerate(table[header_idx]):
        h = _norm(heading)
        if not h:
            continue
        for key, names in ALIASES.items():
            if h in names and key not in column_of:
                column_of[key] = i
                break
        else:
            unknown.append(heading.strip())
    if "url" not in column_of:
        found = ", ".join(c.strip() for c in table[header_idx] if c.strip()) or "(none)"
        raise InputError(
            f"Could not find a URL column in {path.name}. Columns found: {found}. "
            f"Name the column 'URL' (or 'Endpoint' / 'API URL').")

    rows: list[Row] = []
    for offset, raw in enumerate(table[header_idx + 1:], start=header_idx + 2):
        if not any(c.strip() for c in raw):
            continue
        cells = {k: (raw[i] if i < len(raw) else "") for k, i in column_of.items()}
        rows.append(_build_row(len(rows) + 1, offset, cells))
    return Sheet(rows=rows, ignored_columns=unknown)


def load_records(records: list) -> Sheet:
    """Rows from in-memory records keyed by CANONICAL names (Studio's grid, Postman / cURL imports)."""
    rows: list[Row] = []
    for i, rec in enumerate(records, start=1):
        cells = {k: str(rec.get(k, "") if rec.get(k) is not None else "") for k in CANONICAL}
        if not any(v.strip() for v in cells.values()):
            continue
        rows.append(_build_row(len(rows) + 1, i + 1, cells))
    return Sheet(rows=rows, ignored_columns=[])


def sheet_records(path: str | Path) -> list:
    """The raw cells of a sheet as CANONICAL-keyed dicts (what Studio shows in its grid)."""
    table = read_table(Path(path))
    header_idx = next((i for i, r in enumerate(table) if any(c.strip() for c in r)), None)
    if header_idx is None:
        return []
    column_of: dict = {}
    for i, heading in enumerate(table[header_idx]):
        h = _norm(heading)
        for key, names in ALIASES.items():
            if h in names and key not in column_of:
                column_of[key] = i
                break
    if "url" not in column_of:
        found = ", ".join(c.strip() for c in table[header_idx] if c.strip()) or "(none)"
        raise InputError(f"Could not find a URL column. Columns found: {found}. Name the column 'URL'.")
    out = []
    for raw in table[header_idx + 1:]:
        if not any(c.strip() for c in raw):
            continue
        out.append({k: (raw[i] if i < len(raw) else "") for k, i in column_of.items()})
    return out
